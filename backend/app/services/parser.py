import io
import os
import re
from typing import Dict, Any, Optional

try:
    # pyrefly: ignore [missing-import]
    import pymupdf as fitz  # PyMuPDF
except ImportError:
    try:
        # pyrefly: ignore [missing-import]
        import fitz
    except ImportError:
        fitz = None

try:
    # pyrefly: ignore [missing-import]
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None


SECTION_HEADERS = {
    'skills': [
        r'skills', r'technical skills', r'core competencies', r'key skills',
        r'technologies', r'programming languages', r'tools & technologies',
        r'skill set', r'areas of expertise', r'technical proficiencies'
    ],
    'certifications': [
        r'certifications', r'certificates', r'licenses & certifications',
        r'professional certifications', r'courses & certifications',
        r'accreditations', r'training & certifications'
    ],
    'experience': [
        r'experience', r'work experience', r'employment history', r'professional experience',
        r'work history', r'career history', r'internships', r'relevant experience'
    ],
    'education': [
        r'education', r'academic background', r'qualifications', r'educational qualifications',
        r'academic history', r'degrees', r'certifications & education'
    ],
    'projects': [
        r'projects', r'personal projects', r'academic projects', r'key projects',
        r'portfolio', r'open source contributions'
    ]
}


def extract_text_from_pdf(file_bytes: bytes) -> str:
    # 1. Try PyMuPDF (fitz) - high fidelity for multi-column and graphical CVs
    if fitz is not None:
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            pages_text = [page.get_text("text") for page in doc if page.get_text("text").strip()]
            if pages_text:
                return "\n".join(pages_text)
        except Exception:
            pass

    # 2. Try pdfplumber
    if pdfplumber is not None:
        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                pages_text = [page.extract_text() for page in pdf.pages if page.extract_text()]
                if pages_text:
                    return "\n".join(pages_text)
        except Exception:
            pass

    # 3. Fallback to pypdf
    if pypdf is not None:
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            pages_text = [page.extract_text() for page in reader.pages if page.extract_text()]
            if pages_text:
                return "\n".join(pages_text)
        except Exception:
            pass

    return ""


def extract_text_from_docx(file_bytes: bytes) -> str:
    if docx is None:
        return ""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return "\n".join(paragraphs)
    except Exception:
        return ""


def extract_text_from_txt(file_bytes: bytes) -> str:
    for encoding in ['utf-8', 'latin-1', 'cp1252', 'ascii']:
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode('utf-8', errors='ignore')


def extract_candidate_email(text: str) -> Optional[str]:
    email_pattern = r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+'
    match = re.search(email_pattern, text)
    return match.group(0).strip() if match else None


def extract_candidate_phone(text: str) -> Optional[str]:
    phone_pattern = r'(\+?\d{1,3}[-.\s]?)?\(?\d{3,5}\)?[-.\s]?\d{3,5}[-.\s]?\d{3,5}'
    matches = re.findall(phone_pattern, text)
    if matches:
        full_matches = re.finditer(r'(\+?\d{1,3}[-.\s]?)?\(?\d{3,5}\)?[-.\s]?\d{3,5}[-.\s]?\d{3,5}', text)
        for m in full_matches:
            digits = re.sub(r'\D', '', m.group(0))
            if 7 <= len(digits) <= 15:
                return m.group(0).strip()
    return None


def clean_spaced_letters(s: str) -> str:
    """Handle spaced-out text like 'S R I R A A M  V E N K AT E S A N' -> 'SRIRAAM VENKATESAN'."""
    words = [w.strip() for w in re.split(r'\s{2,}', s) if w.strip()]
    if len(words) > 1:
        condensed = [w.replace(' ', '') for w in words]
        if all(len(c) >= 2 and c.isalpha() for c in condensed):
            return ' '.join(condensed)
    if re.match(r'^(?:[A-Za-z]\s+)+[A-Za-z]$', s.strip()):
        return s.replace(' ', '')
    return s.strip()


def is_valid_name_candidate(line: str) -> bool:
    cleaned = clean_spaced_letters(line)
    lower = cleaned.lower()
    
    # Exclude common sentence/resume action words
    verb_stops = [
        'built', 'developed', 'managed', 'created', 'designed', 'organized', 'engineered',
        'implemented', 'spearheaded', 'assisted', 'guided', 'handled', 'coordinated',
        'leading', 'working', 'pursuing', 'student', 'graduate', 'candidate',
        'curriculum', 'resume', 'profile', 'summary', 'contact', 'address', 'experience',
        'education', 'skills', 'objective', 'projects', 'qualification', 'awards',
        'sponsors', 'attendees', 'logistics', 'guidance', 'platform', 'technology',
        'institute', 'university', 'college', 'school', 'chennai', 'india', 'fincard', 'agridirect'
    ]
    if any(re.search(r'\b' + re.escape(v) + r'\b', lower) for v in verb_stops):
        return False
        
    if any(char.isdigit() for char in cleaned) or '@' in cleaned or 'http' in lower or 'www' in lower or 'linkedin' in lower or 'github' in lower:
        return False

    alpha_only = re.sub(r'[^a-zA-Z\s]', '', cleaned).strip()
    words = alpha_only.split()
    if 1 <= len(words) <= 4 and all(len(w) >= 2 and w.isalpha() for w in words) and len(alpha_only) >= 3:
        return True
    return False


def extract_candidate_name(text: str, filename: str = '') -> Optional[str]:
    lines = [line.strip() for line in text.split('\n') if line.strip()]
    if not lines:
        return None

    # Strategy 1: Find candidate email and check lines immediately preceding it (within 4 lines)
    email_idx = -1
    for i, line in enumerate(lines):
        if '@' in line and '.' in line:
            email_idx = i
            break
            
    if email_idx != -1:
        for j in range(email_idx - 1, max(-1, email_idx - 5), -1):
            line = lines[j]
            if is_valid_name_candidate(line):
                cleaned = clean_spaced_letters(line)
                alpha_only = re.sub(r'[^a-zA-Z\s]', '', cleaned).strip()
                if alpha_only:
                    return ' '.join(w.capitalize() for w in alpha_only.split())

    # Strategy 2: If filename is candidate's actual name (not generic like resume.pdf, cv.pdf)
    if filename:
        base = filename.rsplit('.', 1)[0].replace('_', ' ').replace('-', ' ').strip()
        base_clean = re.sub(r'\b(resume|cv|curriculum|vitae|profile|document|biodata|sample|template)\b', '', base, flags=re.IGNORECASE).strip()
        base_alpha = re.sub(r'[^a-zA-Z\s]', '', base_clean).strip()
        words = base_alpha.split()
        if 1 <= len(words) <= 4 and all(len(w) >= 2 and w.isalpha() for w in words):
            return ' '.join(w.capitalize() for w in words)

    # Strategy 3: Scan first 12 lines for a clean person name
    for line in lines[:12]:
        if is_valid_name_candidate(line):
            cleaned = clean_spaced_letters(line)
            alpha_only = re.sub(r'[^a-zA-Z\s]', '', cleaned).strip()
            if alpha_only:
                return ' '.join(w.capitalize() for w in alpha_only.split())

    return None


def parse_structural_sections(text: str) -> Dict[str, str]:
    sections: Dict[str, str] = {
        'skills': '',
        'certifications': '',
        'experience': '',
        'education': '',
        'projects': '',
        'other': ''
    }

    lines = text.split('\n')
    current_section = 'other'
    section_buffers: Dict[str, list] = {k: [] for k in sections.keys()}

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        lower_line = stripped.lower().strip(':').strip('-').strip('#').strip()
        matched_section = None

        for sec_name, patterns in SECTION_HEADERS.items():
            for pat in patterns:
                if re.fullmatch(pat, lower_line, re.IGNORECASE) or (len(lower_line.split()) <= 4 and re.search(r'\b' + pat + r'\b', lower_line)):
                    matched_section = sec_name
                    break
            if matched_section:
                break

        if matched_section:
            current_section = matched_section
        else:
            section_buffers[current_section].append(stripped)

    for sec, buffer in section_buffers.items():
        sections[sec] = "\n".join(buffer).strip()

    return sections


def parse_resume(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    ext = filename.lower().split('.')[-1] if '.' in filename else ''
    
    if ext == 'pdf':
        raw_text = extract_text_from_pdf(file_bytes)
    elif ext in ['docx', 'doc']:
        raw_text = extract_text_from_docx(file_bytes)
    elif ext == 'txt':
        raw_text = extract_text_from_txt(file_bytes)
    else:
        raw_text = extract_text_from_txt(file_bytes)

    if not raw_text.strip():
        raw_text = extract_text_from_txt(file_bytes)

    candidate_name = extract_candidate_name(raw_text, filename)
    candidate_email = extract_candidate_email(raw_text)
    candidate_phone = extract_candidate_phone(raw_text)
    sections = parse_structural_sections(raw_text)

    return {
        'filename': filename,
        'file_type': ext or 'unknown',
        'raw_text': raw_text,
        'candidate_name': candidate_name,
        'candidate_email': candidate_email,
        'candidate_phone': candidate_phone,
        'sections': sections
    }
